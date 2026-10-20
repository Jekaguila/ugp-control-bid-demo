"""Verifica que el cálculo independiente en pandas coincide con las fórmulas del libro de Excel.

Ejecutar desde la carpeta dashboard/:  python -m pytest -q
Requiere que el libro tenga valores calculados (guardado por Excel/LibreOffice).
"""
import sys
from pathlib import Path
import numpy as np
import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ugp.cargar_datos import cargar, RUTA_POR_DEFECTO  # noqa: E402
from ugp import indicadores as ind  # noqa: E402


@pytest.fixture(scope="module")
def datos():
    return cargar(RUTA_POR_DEFECTO)


@pytest.fixture(scope="module")
def libro():
    return openpyxl.load_workbook(RUTA_POR_DEFECTO, data_only=True)


def _tablero(libro):
    ws = libro["Tablero"]
    return {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(6, 22)}


def test_controles_del_libro_en_ok(libro):
    assert libro["Controles"]["C22"].value == "TODOS LOS CONTROLES OK"


def test_montos_del_programa(datos):
    p = datos["param"]
    assert p["Total_Prestamo"] == 100_000_000 and p["Total_Donacion"] == 4_000_000
    assert p["Total_Programa"] == 104_000_000
    pr = datos["presupuesto"]
    assert pr[pr["Comp."].astype(str).str.match(r"^C\d$")]["Total asignado"].sum() == 104_000_000


@pytest.mark.parametrize("etiqueta,clave,tol", [
    ("Ejecutado acumulado (USD)", "ejecutado", 0.5),
    ("% ejecutado sobre el total del Programa", "pct_ejecutado", 1e-6),
    ("Avance físico del POA (real vs. planificado a la fecha)", "avance_fisico", 1e-6),
    ("Avance financiero del POA (ejecutado / presupuesto POA)", "avance_financiero", 1e-6),
    ("SPI — índice de desempeño del cronograma", "SPI", 1e-6),
    ("CPI — índice de desempeño del costo", "CPI", 1e-6),
    ("Total desembolsado por el BID (USD)", "total_desembolsado", 0.5),
    ("% de anticipos justificado", "pct_justificado", 1e-6),
    ("Riesgos altos abiertos", "riesgos_altos", 0),
    ("Obligaciones normativas vencidas", "obligaciones_vencidas", 0),
    ("% de cumplimiento de obligaciones exigibles", "pct_cumplimiento", 1e-6),
])
def test_kpis_python_vs_excel(datos, libro, etiqueta, clave, tol):
    xl = _tablero(libro)[etiqueta]
    py = ind.kpis(datos)[clave]
    assert abs(float(xl) - float(py)) <= tol, f"{etiqueta}: Excel={xl} Python={py}"


def test_evm_por_actividad(datos):
    poa = ind.poa_evm(datos)
    xl = datos["poa"].set_index("ID")
    for _, r in poa.iterrows():
        assert r["EV"] == pytest.approx(xl.loc[r["ID"], "EV USD"], abs=0.5)
        assert r["PV"] == pytest.approx(xl.loc[r["ID"], "PV USD"], abs=0.5)
        assert r["AC"] == pytest.approx(xl.loc[r["ID"], "Ejecutado USD (AC)"], abs=0.5)


def test_semaforo_por_actividad(datos):
    poa = ind.poa_evm(datos)
    xl = datos["poa"].set_index("ID")["Semáforo"]
    for _, r in poa.iterrows():
        assert r["sem"] == xl.loc[r["ID"]]


def test_indice_focalizacion(datos):
    f = ind.focalizacion(datos).set_index("Departamento")
    xl = datos["focal"].set_index("Departamento")
    xl = xl[xl["ISO 3166-2"].astype(str).str.startswith("SV-")]
    assert len(f) == 14
    col_inv = next(c for c in xl.columns if c.startswith("Inversión ejecutada"))
    for dep in f.index:
        assert f.loc[dep, "indice"] == pytest.approx(xl.loc[dep, "Índice de focalización"], abs=1e-9)
        assert f.loc[dep, "prioridad"] == xl.loc[dep, "Prioridad"]
        assert f.loc[dep, "inversion"] == pytest.approx(xl.loc[dep, col_inv], abs=0.5)


def test_integridad_del_libro_de_gastos(datos):
    g = datos["gastos"]
    assert g["Monto USD"].ge(0).all()
    assert g["ID gasto"].is_unique
    assert set(g["Actividad POA"]) <= set(datos["poa"]["ID"])
    assert set(g["Contrato"].dropna()) - {"N/A"} <= set(datos["contratos"]["ID contrato"])
    ini, corte = datos["param"]["Fecha_Inicio"], datos["param"]["Fecha_Corte"]
    assert g["Fecha"].map(lambda x: ini <= x <= corte).all()


def test_pagos_directos_conciliados(datos):
    a = ind.anticipos(datos)
    g = datos["gastos"]
    pd_libro = g.loc[(g["Mecanismo"] == "Pago directo") & (g["Estado de pago"] == "Pagado"), "Monto USD"].sum()
    assert a["pagos_directos"] == pytest.approx(pd_libro, abs=0.5)
    assert a["pagado_con_anticipo"] <= a["anticipos"]
