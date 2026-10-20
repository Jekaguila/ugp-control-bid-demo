"""Hojas base: Parametros y Listas."""
from openpyxl.styles import Alignment
from xl_base import *
from datos_demo import *  # noqa: F401,F403

LISTAS = {
    "Componente": [c[0] for c in COMPONENTES],
    "Fuente": ["Préstamo", "Donación"],
    "Categoria_gasto": ["Consultoría", "Servicios de no consultoría", "Bienes", "Obras",
                        "Transferencias a beneficiarios", "Costos operativos"],
    "Mecanismo": ["Anticipo de fondos", "Reembolso de gastos", "Pago directo"],
    "Estado_pago": ["Pagado", "Por pagar"],
    "Justificado": ["Sí", "No", "N/A"],
    "Departamento": ["Nacional"] + [d[0] for d in DEPTOS],
    "Metodo_adq": ["LPI", "LPN", "Comparación de precios", "Contratación directa",
                   "SBCC", "SBC", "SBMC", "3CV"],
    "Estado_adq": ["Planificado", "Preparación", "Publicado", "Evaluación", "Contratado", "Cerrado"],
    "Categoria_adq": ["Bienes", "Obras", "Servicios de no consultoría", "Consultoría (firma)", "Consultor individual"],
    "Revision": ["Ex ante", "Ex post"],
    "Estado_cumpl": ["Pendiente", "En proceso", "Cumplida"],
    "Estado_compromiso": ["Pendiente", "En proceso", "Cumplido"],
    "Estado_NO": ["Pendiente", "Otorgada", "Con observaciones"],
    "Estado_riesgo": ["Abierto", "En mitigación", "Cerrado"],
    "Categoria_riesgo": ["Técnico", "Fiduciario", "Institucional", "Ambiental y social", "Cumplimiento", "Externo"],
    "Respuesta_riesgo": ["Mitigar", "Evitar", "Transferir", "Aceptar"],
    "Actitud": ["Líder", "A favor", "Neutral", "En contra"],
    "Conformidad": ["Conforme", "Con observaciones", "Pendiente"],
    "Escala": [1, 2, 3, 4, 5],
    "Tipo_meta": ["Acumulativa", "Nivel"],
    "Responsable": [R_CG, R_FIN, R_ADQ, R_PLAN, R_SOC, "Gerencia de Cooperación Institucional", "Dirección General"],
    "Instrumento": ["Contrato de Préstamo 6153/OC-ES", "Convenio GR 6154/GR-ES", "MOP",
                    "Políticas BID", "Normativa nacional", "TDR / UGP"],
    "Estado_audit": ["Planificada", "En ejecución", "Informe emitido", "Cerrada"],
    "Estado_hallazgo": ["Abierta", "En proceso", "Cerrada"],
}
REF = {}   # nombre de lista -> rango absoluto, se llena en hoja_listas


def hoja_listas(wb):
    ws = wb.create_sheet("Listas")
    titulo(ws, "Listas de valores (fuentes de las listas desplegables)")
    for j, (k, vals) in enumerate(LISTAS.items(), 1):
        ws.cell(HDR, j, k).font = f(True, "FFFFFF")
        ws.cell(HDR, j).fill = C_ENCAB
        ws.column_dimensions[L(j)].width = max(16, len(k) + 4)
        for i, v in enumerate(vals):
            ws.cell(HDR + 1 + i, j, v).font = f(False, "0000FF")
        REF[k] = f"Listas!${L(j)}${HDR + 1}:${L(j)}${HDR + len(vals)}"
    nota(ws, 2 + 1, "Para agregar un valor, escríbalo debajo de la última fila de su lista y amplíe el rango en Datos > Validación de datos.")
    return ws


def hoja_parametros(wb):
    ws = wb.create_sheet("Parametros")
    titulo(ws, "Parámetros del Programa y del sistema de control",
           "Celdas azules = entradas editables. Las celdas con nombre (p. ej. Fecha_Corte) alimentan todas las hojas.")
    cols = [dict(h="Grupo", w=22, tipo="lbl"), dict(h="Parámetro", w=52, tipo="lbl", wrap=True),
            dict(h="Valor", w=34, al="right"), dict(h="Unidad", w=14, tipo="lbl"),
            dict(h="Nombre en fórmulas", w=22, tipo="lbl"), dict(h="Fuente / nota", w=70, wrap=True, tipo="lbl")]
    P = [
        ("Programa", "Programa", "Programa de Integración Académica y Comunitaria para Jóvenes en El Salvador", "", "", "TDR, carátula", None),
        ("Programa", "Operaciones BID", "ES-L1175 y ES-J0003", "", "", "TDR; CCLIP ES-O0023", None),
        ("Programa", "Contrato de Préstamo", "6153/OC-ES", "", "", "TDR (el texto del usuario decía «6153 OCS»; el TDR dice 6153/OC-ES)", None),
        ("Programa", "Convenio de Financiamiento No Reembolsable", "6154/GR-ES", "", "", "TDR", None),
        ("Programa", "Resolución de aprobación BID", "DE-43/26 (24-jun-2026)", "", "", "TDR §I", None),
        ("Programa", "Organismo Ejecutor", "Dirección de Integración", "", "", "TDR; la UGP está adscrita a la Gerencia de Cooperación Institucional", None),
        ("Programa", "Monto del préstamo", 100_000_000, "USD", "Total_Prestamo", "TDR §I", USD),
        ("Programa", "Monto de la donación", 4_000_000, "USD", "Total_Donacion", "TDR §I", USD),
        ("Programa", "Total financiamiento BID", "=C11+C12", "USD", "Total_Programa", "Calculado", USD),
        ("Programa", "Duración", 5, "años", "Duracion_Anios", "TDR §I", "0"),
        ("Calendario", "Fecha de inicio de ejecución (SUPUESTO)", INICIO, "fecha", "Fecha_Inicio", "Supuesto de la demo: la fecha real depende de la entrada en vigencia", FECHA),
        ("Calendario", "Fecha de corte del reporte (SIMULADA)", CORTE, "fecha", "Fecha_Corte", "Cambie esta fecha y todo el sistema se recalcula", FECHA),
        ("Calendario", "Año del programa a la fecha de corte", "=INT((Fecha_Corte-Fecha_Inicio)/365.25)+1", "año", "Anio_Programa", "Calculado", "0"),
        ("Calendario", "Fracción transcurrida del año del programa", "=(Fecha_Corte-Fecha_Inicio)/365.25-(Anio_Programa-1)", "%", "Frac_Anio", "Calculado", PCT),
        ("Calendario", "Inicio de la consultoría del Coordinador General (SUPUESTO)", INICIO_CONSULTORIA, "fecha", "Inicio_Consultoria", "Supuesto: plazo de 12 meses desde la firma (TDR §VIII)", FECHA),
        ("Consultoría", "Honorario mensual del Coordinador General", 4_200, "USD", "Pago_Mensual", "TDR §XII (US$50,400 en 12 pagos)", USD),
        ("Consultoría", "Plazo de entrega del informe mensual", 3, "días hábiles", "Plazo_Informe", "TDR §VII: tres días hábiles posteriores al cierre del mes", "0"),
        ("Semáforos", "Indicadores: umbral verde (cumplimiento ≥)", 0.90, "ratio", "Umbral_Verde", "Criterio de gestión propuesto", PCT),
        ("Semáforos", "Indicadores: umbral amarillo (cumplimiento ≥)", 0.70, "ratio", "Umbral_Amarillo", "Criterio de gestión propuesto", PCT),
        ("Semáforos", "SPI/CPI: umbral verde (≥)", 0.95, "ratio", "EVM_Verde", "Criterio de gestión propuesto", "0.00"),
        ("Semáforos", "SPI/CPI: umbral amarillo (≥)", 0.85, "ratio", "EVM_Amarillo", "Criterio de gestión propuesto", "0.00"),
        ("Semáforos", "Riesgo ALTO si probabilidad×impacto ≥", 15, "puntos", "Riesgo_Alto", "Escala 1-5 × 1-5", "0"),
        ("Semáforos", "Riesgo MEDIO si probabilidad×impacto ≥", 8, "puntos", "Riesgo_Medio", "Escala 1-5 × 1-5", "0"),
        ("Semáforos", "Alerta de vencimiento (días antes)", 30, "días", "Dias_Alerta", "Criterio de gestión propuesto", "0"),
        ("Fiduciario", "Justificación mínima de anticipos para pedir uno nuevo", 0.80, "%", "Pct_Justif_Min", "A VALIDAR con el Contrato de Préstamo / guía de gestión financiera del BID", PCT),
        ("Fiduciario", "Plazo de referencia de trámite de No Objeción", 15, "días", "Plazo_NO", "Referencia INTERNA de la UGP, no un plazo del BID", "0"),
        ("Fiduciario", "Plazo de entrega de EEFF auditados tras el cierre", 120, "días", "Plazo_EEFF", "A VALIDAR con el Contrato de Préstamo", "0"),
        ("Focalización", "Peso de la vulnerabilidad en el índice", 0.50, "peso", "W_Vuln", "Criterio propuesto (los pesos suman 1)", "0.00"),
        ("Focalización", "Peso de la población objetivo en el índice", 0.20, "peso", "W_Pob", "Criterio propuesto", "0.00"),
        ("Focalización", "Peso de la brecha de cobertura en el índice", 0.30, "peso", "W_Brecha", "Criterio propuesto", "0.00"),
        ("Focalización", "Prioridad ALTA si índice ≥", 0.60, "índice", "Prior_Alta", "Criterio propuesto", "0.00"),
        ("Focalización", "Prioridad MEDIA si índice ≥", 0.45, "índice", "Prior_Media", "Criterio propuesto", "0.00"),
        ("Focalización", "Meta del año 1 como % de la meta a 5 años", 0.10, "%", "Pct_Anio1_Benef", "Supuesto de la demo", PCT),
        ("Focalización", "Meta de beneficiarios a 5 años (SIMULADA)", META_BENEF_5A, "personas", "Meta_Benef_5A", "Dato sintético", "#,##0"),
    ]
    filas = [[g, p, v, u, n, o] for g, p, v, u, n, o, _ in P]
    tabla(ws, cols, filas, filtro=False, altura=26)
    for i, (g, p, v, u, n, o, fm) in enumerate(P):
        r = HDR + 1 + i
        c = ws.cell(r, 3)
        c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        if fm:
            c.number_format = fm
        if isinstance(v, str) and v.startswith("="):
            c.fill, c.font = C_CALC, f()
        if n:
            nombre(wb, n, "Parametros", f"$C${r}")
    ws.freeze_panes = "A5"
    return ws
