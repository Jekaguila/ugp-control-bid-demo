"""Utilidades de estilo y escritura para los libros de Excel (openpyxl)."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.defined_name import DefinedName

FUENTE = "Arial"
AZUL = "1F3864"
C_ENCAB = PatternFill("solid", fgColor=AZUL)
C_CALC = PatternFill("solid", fgColor="F2F2F2")
C_SUB = PatternFill("solid", fgColor="D9E2F3")
C_AVISO = PatternFill("solid", fgColor="FFF2CC")
VERDE, AMARILLO, ROJO = "C6EFCE", "FFEB9C", "FFC7CE"
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
FIRST = 5          # primera fila de datos en todas las hojas de tabla
HDR = 4            # fila de encabezados
USD = '#,##0;[Red]-#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
FECHA = "dd/mm/yyyy"

AVISO = ("DATOS SIMULADOS con fines de demostración metodológica — no son cifras reales del BID "
         "ni de la Dirección de Integración.")


def f(bold=False, color="000000", size=11, italic=False):
    return Font(name=FUENTE, bold=bold, color=color, size=size, italic=italic)


def titulo(ws, texto, sub=None, ancho=10):
    ws["A1"] = texto
    ws["A1"].font = f(True, AZUL, 16)
    ws["A2"] = AVISO
    ws["A2"].font = f(False, "C00000", 10, True)
    if sub:
        ws["A3"] = sub
        ws["A3"].font = f(False, "595959", 10)
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 24


def nombre(wb, nom, hoja, celda):
    wb.defined_names[nom] = DefinedName(nom, attr_text=f"'{hoja}'!{celda}")


def tabla(ws, cols, filas, hdr=HDR, filtro=True, altura=34):
    """cols: lista de dicts {h, w, tipo('in'|'calc'), fmt}. filas: lista de listas (valores o fórmulas)."""
    for j, c in enumerate(cols, 1):
        cel = ws.cell(hdr, j, c["h"])
        cel.font = f(True, "FFFFFF")
        cel.fill = C_ENCAB
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = BORDE
        ws.column_dimensions[L(j)].width = c["w"]
    ws.row_dimensions[hdr].height = altura
    for i, fila in enumerate(filas):
        r = hdr + 1 + i
        for j, (c, v) in enumerate(zip(cols, fila), 1):
            cel = ws.cell(r, j, v)
            calc = c.get("tipo") == "calc"
            neutro = calc or c.get("tipo") == "lbl"
            cel.font = f(False, "000000" if neutro else "0000FF")
            if calc:
                cel.fill = C_CALC
            cel.border = BORDE
            if c.get("fmt"):
                cel.number_format = c["fmt"]
            cel.alignment = Alignment(vertical="top", wrap_text=c.get("wrap", False),
                                      horizontal=c.get("al"))
    ws.freeze_panes = ws.cell(hdr + 1, c_freeze(cols))
    if filtro and filas:
        ws.auto_filter.ref = f"A{hdr}:{L(len(cols))}{hdr + len(filas)}"


def c_freeze(cols):
    return 3 if len(cols) > 6 else 2


def lista(ws, rango, fuente, permitir_vacio=True):
    dv = DataValidation(type="list", formula1=fuente, allow_blank=permitir_vacio)
    dv.error, dv.errorTitle = "Valor no válido", "Seleccione un valor de la lista"
    ws.add_data_validation(dv)
    dv.add(rango)


def semaforo(ws, rango, verde, amarillo, rojo):
    """Reglas por texto exacto en la celda."""
    for txt, color in ((verde, VERDE), (amarillo, AMARILLO), (rojo, ROJO)):
        if txt:
            ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=[f'"{txt}"'],
                                                            fill=PatternFill("solid", bgColor=color, fgColor=color)))


def semaforo_ratio(ws, rango, v="Umbral_Verde", a="Umbral_Amarillo"):
    """Colorea valores numéricos (ratios) usando umbrales con nombre."""
    celda = rango.split(":")[0].replace("$", "")
    for expr, color in ((f"AND(ISNUMBER({celda}),{celda}>={v})", VERDE),
                        (f"AND(ISNUMBER({celda}),{celda}<{v},{celda}>={a})", AMARILLO),
                        (f"AND(ISNUMBER({celda}),{celda}<{a})", ROJO)):
        ws.conditional_formatting.add(rango, FormulaRule(formula=[expr], fill=PatternFill("solid", bgColor=color, fgColor=color)))


def nota(ws, fila, texto, col=1):
    c = ws.cell(fila, col, texto)
    c.font = f(False, "595959", 10, True)
    return c


def rng(hoja, col, ultima=500, primera=FIRST):
    return f"{hoja}!${col}${primera}:${col}${ultima}"
